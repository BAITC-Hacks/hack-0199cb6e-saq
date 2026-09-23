"""Small synthetic workbooks with the partner headers from SPEC_DATA."""

from __future__ import annotations

from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl.comments import Comment

from ekt_adapters import load_iek, load_se
from ekt_adapters.cli import main
from ekt_adapters.xlsx_safe import open_workbook
from ekt_engine import run


def _book(path: Path, title: str, rows: list[list[object]], *, header_row: int = 1):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = title
    for _ in range(header_row - 1):
        sheet.append([])
    for row in rows:
        sheet.append(row)
    book.save(path)
    return path


@pytest.fixture
def partner_dir(tmp_path: Path) -> Path:
    _book(
        tmp_path / "se-sales.xlsx", "Лист_1",
        [
            ["Номенклатура", "Номенклатура.Код", "Артикул", "Кратность",
             "июль 2026", "авг. 2026", "сент. 2026", "Итого"],
            ["Количество"],
            ["Синтетика SE", "030200192_", "A-1", 2, 100, 120, 20, "=SUM(E3:G3)"],
            ["Итого"],
        ],
    )
    _book(
        tmp_path / "se-stock.xlsx", "Лист_1",
        [
            ["№", "Номенклатура", "Номенклатура.Код", "Ед.изм",
             "авг. 2026", "сент. 2026"],
            [], [],
            [1, "Синтетика SE", "030200192_", "шт", 100, 200],
        ],
    )
    _book(
        tmp_path / "se-moq.xlsx", "Лист_1",
        [
            ["№", "Номенклатура", "Номенклатура.Код", "Артикул", "Кратность"],
            [], [1, "Синтетика SE", "030200192_", "A-1", 10],
        ],
    )
    se_master = _book(
        tmp_path / "se-master.xlsx", "TDSheet",
        [
            ["№", "Артикул поставщика", "Код 1с", "Наименование",
             "Категория 2026", "Свободный остаток", "Зарезервировано",
             "Розничный склад", "Витрина", "Остаток ТЗ", "РЦ ЕКТ  Рыскулова",
             "СЭ в пути 24.09"],
            [1, "A-1", "030200192_", "Синтетика SE", "1", 90, 0, 0, 0, 0, 0, 10],
        ],
        header_row=2,
    )
    book = openpyxl.load_workbook(se_master)
    book.active["L3"].comment = Comment("Поставка 20.10", "synthetic-author")
    book.save(se_master)

    _book(
        tmp_path / "iek-sales.xlsx", "Лист_1",
        [
            ["Номенклатура", "Номенклатура.Код", "июль 2026",
             "авг. 2026", "сент. 2026", "Итого"],
            ["Количество"],
            ["Синтетика IEK", "щт23054819", 30, 40, 5, 75],
            ["Итого"],
        ],
    )
    _book(
        tmp_path / "iek-stock.xlsx", "Лист_1",
        [
            ["Номенклатура", "Ед.", "Номенклатура.Код",
             "авг. 2026", "сент. 2026", "Итого"],
            ["Количество"], ["нач. остаток"],
            ["Синтетика IEK", "м", "щт23054819", "40,5", 30, 70.5],
            ["Итого"],
        ],
    )
    _book(
        tmp_path / "iek-moq.xlsx", "Лист7",
        [
            ["№", "Код 1с", "Артикул поставщика", "Наименование",
             "Мин. разр. к отгр."],
            [1, "щт23054819", "B-2", "Синтетика IEK", 5],
        ],
    )
    _book(
        tmp_path / "iek-path.xlsx", "Лист4",
        [
            ["Код 1с", "Артикул ИЭК", " Наименование",
             "РФ  УТ-7583 от 31 августа 2026\xa0г. (поступление до 10.10.2026)"],
            ["щт23054819", "B-2", "Синтетика IEK", 15],
        ],
    )
    _book(
        tmp_path / "se-lines.xlsx", "Лист_1",
        [
            ["Дата", "Номер", "Документ", "Код", "Номенклатура",
             "Ед.", "Склад", "Количество"],
            ["15.08.2026 9:00:00", 12, "Расходная накладная", "030200192_",
             "Синтетика SE", "шт", "Алматы", 120],
            ["22.09.2026 9:00:00", 13, "Расходная накладная", "030200192_",
             "Синтетика SE", "шт", "Алматы", 20],
            ["22.09.2026 9:00:00", 14, "Заказ покупателя", "030200192_",
             "Синтетика SE", "шт", "Алматы", 999],
            ["01.12.2024 9:00:00", 15, "Расходная накладная", "030200192_",
             "Синтетика SE", "шт", "Алматы", 1],
            ["Итого"],
        ],
    )
    _book(
        tmp_path / "iek-lines.xlsx", "Лист_1",
        [
            ["Дата", "Номер", "Документ", "Код", "Номенклатура",
             "Ед.", "Склад", "Количество"],
            ["15.08.2026 9:00:00", 4, "Расходная накладная", "щт23054819",
             "Синтетика IEK", "м", "Алматы", 40],
            ["22.09.2026 9:00:00", 5, "Расходная накладная", "щт23054819",
             "Синтетика IEK", "м", "Алматы", 5],
            ["Итого"],
        ],
    )
    return tmp_path


def test_se_loads_pipeline_tables(partner_dir: Path):
    result = load_se(partner_dir)
    assert len(result.skus) == 1
    assert result.skus[0]["sku_id"] == "030200192_"
    assert result.skus[0]["order_multiple"] == 10
    assert result.skus[0]["abc"] == "A"
    assert [row["month"] for row in result.sales_monthly] == ["2026-07", "2026-08"]
    assert len(result.sales_lines) == 2
    assert result.sales_lines[-1]["doc_key"] == "2026-13"
    assert result.stock_current[0]["free"] == 90
    assert result.inbound[0]["eta"].isoformat() == "2026-10-20"
    assert result.inbound[0]["source"] == "comment"
    workbook = open_workbook(partner_dir / "se-sales.xlsx")
    try:
        assert workbook.active["H3"].value is None  # data_only=True: no formula execution
    finally:
        workbook.close()
    assert run(result).by_sku("030200192_")["qty_recommended"] >= 0


def test_iek_loads_pipeline_tables(partner_dir: Path):
    result = load_iek(partner_dir)
    assert len(result.skus) == 1
    assert result.skus[0]["sku_id"] == "щт23054819"
    assert result.skus[0]["order_multiple"] == 5
    assert result.skus[0]["uom_sales"] == "м"
    assert [row["month"] for row in result.sales_monthly] == ["2026-07", "2026-08"]
    assert result.stock_opening[-2]["qty"] == 40.5
    assert result.stock_current[0]["free"] == 25
    assert result.stock_current[0]["is_estimate"] is True
    assert result.inbound[0]["channel_id"] == "IEK-RF"
    assert result.inbound[0]["eta"].isoformat() == "2026-10-10"
    assert run(result).by_sku("щт23054819")["qty_recommended"] >= 0


def test_iek_moq_na_uses_one(partner_dir: Path):
    source = partner_dir / "iek-moq.xlsx"
    book = openpyxl.load_workbook(source)
    book.active["E2"] = "#N/A"
    book.save(source)
    result = load_iek(partner_dir)
    assert result.skus[0]["order_multiple"] == 1.0
    assert result.skus[0]["min_order_qty"] == 1.0


def test_cli_prints_only_aggregates(partner_dir: Path, capsys):
    assert main(["load", str(partner_dir)]) == 0
    output = capsys.readouterr().out
    assert "SE: sales_monthly=2 sales_lines=2" in output
    assert "IEK: sales_monthly=2 sales_lines=2" in output
    assert "030200192_" not in output
    assert "щт23054819" not in output


def test_rejects_raw_client_column(partner_dir: Path):
    source = partner_dir / "se-lines.xlsx"
    book = openpyxl.load_workbook(source)
    book.active.cell(1, 9, "Контрагент")
    book.save(source)
    with pytest.raises(ValueError, match="CLIENT_ID_RAW"):
        load_se(partner_dir)


def test_rejects_macro_extension(partner_dir: Path):
    with pytest.raises(ValueError, match="Only .xlsx"):
        open_workbook(partner_dir / "not-a-workbook.xlsm")
